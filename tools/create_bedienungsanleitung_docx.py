from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "KitaFenster_Bedienungsanleitung.docx"

GREEN = "195953"
PALE_GREEN = "EAF2F0"
PALE_GOLD = "FAF3DE"
LIGHT_GRAY = "D9D9D9"
MID_GRAY = "6A7774"
INK = "203B3B"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Seite ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def add_step(doc, number, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.15)
    p.paragraph_format.first_line_indent = Cm(-0.15)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(f"{number}  {title}")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(GREEN)
    p.add_run(f"\n{text}")
    return p


def add_tip(doc, lead, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(lead + " ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(GREEN)
    p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.75)
section.bottom_margin = Cm(1.95)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.widow_control = True

title = styles["Title"]
title.font.name = "Aptos Display"
title.font.size = Pt(30)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(6)
title_ppr = title._element.get_or_add_pPr()
title_borders = title_ppr.find(qn("w:pBdr"))
if title_borders is not None:
    title_ppr.remove(title_borders)

for style_name, size, before, after in (
    ("Heading 1", 19, 18, 7),
    ("Heading 2", 13.5, 11, 5),
):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for style_name in ("List Bullet", "List Bullet 2"):
    style = styles[style_name]
    style.font.name = "Aptos"
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor.from_string(INK)
    style.paragraph_format.space_after = Pt(3)

footer = section.footer
footer.distance = Cm(0.65)
fp = footer.paragraphs[0]
fp.add_run("KitaFenster Bedienungsanleitung")
fp.runs[0].font.size = Pt(8.5)
fp.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
fp.add_run("    ")
add_page_number(fp)

# Titelseite
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(24)
r = p.add_run("KITAFENSTER")
r.bold = True
r.font.name = "Aptos"
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string(GREEN)
r.font.letter_spacing = Pt(1.5)

p = doc.add_paragraph(style="Title")
p.add_run("Bedienungsanleitung")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
r = p.add_run("Fotos und Informationen einfach auf den Bilderrahmen anzeigen")
r.font.size = Pt(14)
r.font.color.rgb = RGBColor.from_string(MID_GRAY)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(14)
r = p.add_run("Für wen ist diese Anleitung gedacht")
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = RGBColor.from_string(GREEN)
p.add_run(
    "\nFür alle Mitarbeitenden, die KitaFenster im Alltag bedienen. Technische Vorkenntnisse sind nicht nötig."
)

doc.add_heading("Das Wichtigste in einer Minute", level=1)
add_bullet(
    doc,
    "Fotos kommen in den passenden Gruppenordner auf dem NAS. Die Bilderrahmen übernehmen sie automatisch.",
)
add_bullet(doc, "Meldungen und Dokumente werden in der Webverwaltung unter Inhalte angelegt.")
add_bullet(
    doc, "Unter Bilderrahmen lässt sich festlegen, was ein Rahmen zeigt und wann er eingeschaltet ist."
)
add_bullet(
    doc, "Sofortmeldungen unterbrechen die normale Anzeige. Verwende sie nur für wirklich dringende Hinweise."
)

add_tip(
    doc,
    "Merke",
    "Nach Änderungen kann es einen kurzen Moment dauern, bis ein Bilderrahmen den neuen Stand übernommen hat.",
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
r = p.add_run("Webverwaltung")
r.bold = True
p.add_run("\nhttp://192.168.1.6:8088/")

doc.add_page_break()

# Anmeldung und Orientierung
doc.add_heading("Anmelden und orientieren", level=1)
add_step(
    doc,
    "1",
    "Webseite öffnen",
    "Öffne die Webverwaltung im Browser und melde dich mit deinem Benutzernamen und Passwort an.",
)
add_step(
    doc,
    "2",
    "Bereich auswählen",
    "Die wichtigsten Bereiche findest du im Menü: Gruppen und NAS, Inhalte, Bilderrahmen, Betriebszeiten und Einrichtung.",
)
add_step(
    doc,
    "3",
    "Änderung speichern",
    "Prüfe deine Auswahl und wähle Speichern oder Einstellungen übernehmen. Erst dann wird die Änderung übernommen.",
)

table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Cm(4.7)
table.columns[1].width = Cm(11.5)
mark_header_row(table.rows[0])
headers = ("Bereich", "Dort erledigst du")
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = text
    set_cell_shading(cell, GREEN)
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
for left, right in (
    ("Gruppen und NAS", "Gruppennamen, Gruppenbilder und Fotoordner festlegen"),
    ("Inhalte", "Meldungen und Dokumente anlegen oder ändern"),
    ("Bilderrahmen", "Anzeige, Aufteilung, Reihenfolge und einzelne Geräte verwalten"),
    ("Betriebszeiten", "Gemeinsame Ein- und Ausschaltzeiten festlegen"),
    ("Einrichtung", "Name und Logo der Kita ändern"),
):
    cells = table.add_row().cells
    cells[0].text = left
    cells[1].text = right
    for c in cells:
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(cells[0], PALE_GREEN)
for row in table.rows:
    for cell in row.cells:
        set_cell_margins(cell)
set_table_borders(table)

doc.add_heading("Einmalige Grundeinrichtung", level=1)
add_bullet(doc, "Unter Einrichtung den Namen und das Logo der Kita festlegen.")
add_bullet(
    doc, "Unter Gruppen und NAS die vorhandenen Gruppen umbenennen und den richtigen Fotoordner auswählen."
)
add_bullet(doc, "Unter Bilderrahmen jedem Gerät einen verständlichen Namen und die passende Gruppe zuweisen.")
add_tip(
    doc,
    "Tipp",
    "Gib Gruppen und Bilderrahmen eindeutige Namen, zum Beispiel Marienkäfer oder Eingang Elternbereich.",
)

# Fotos
doc.add_heading("Fotos anzeigen", level=1)
add_step(doc, "1", "Gruppe wählen", "Öffne auf dem NAS den Fotoordner der gewünschten Gruppe.")
add_step(
    doc,
    "2",
    "Fotos kopieren",
    "Kopiere die fertigen Bilder vollständig in diesen Ordner. Geeignet sind JPG, JPEG, PNG, WebP und BMP.",
)
add_step(
    doc,
    "3",
    "Kurz warten",
    "Der Server und die Bilderrahmen gleichen neue Inhalte automatisch ab. Du musst normalerweise nichts weiter tun.",
)
add_tip(doc, "Wichtig", "HEIC-Fotos müssen vorher in JPG oder PNG umgewandelt werden.")
add_tip(
    doc,
    "Beim Leeren eines Ordners",
    "Ein leerer, erreichbarer Gruppenordner bedeutet: Auf diesem Rahmen sollen keine Gruppenfotos mehr gezeigt werden. Wenn der Ordner nicht erreichbar ist, bleibt dagegen der letzte erfolgreiche Stand erhalten.",
)

# Meldungen
doc.add_heading("Meldung veröffentlichen", level=1)
add_step(doc, "1", "Meldung beginnen", "Öffne Inhalte und wähle Meldung erstellen.")
add_step(
    doc,
    "2",
    "Inhalt gestalten",
    "Wähle eine Vorlage oder schreibe einen kurzen Text. Lege Zielgruppe und Farbe fest. Bei Bedarf kannst du ein Bild hinzufügen.",
)
add_step(
    doc,
    "3",
    "Vorschau prüfen",
    "Achte besonders auf Lesbarkeit und Rechtschreibung. Kurze Texte sind auf den Rahmen besser zu lesen.",
)
add_step(doc, "4", "Veröffentlichen", "Lege bei Bedarf ein Ende fest und veröffentliche die Meldung.")
add_tip(
    doc,
    "Sofortmeldung",
    "Diese Meldung erscheint im Vollbild und hat Vorrang vor allen anderen Inhalten, auch außerhalb der Betriebszeiten. Verwende sie nur für dringende Informationen.",
)
add_tip(
    doc,
    "Interne Meldung",
    "Aktiviere Intern, wenn die Meldung nur auf dafür freigegebenen Personalrahmen erscheinen darf.",
)

# Dokumente
doc.add_heading("Dokument anzeigen", level=1)
add_step(
    doc,
    "1",
    "Datei hochladen",
    "Öffne Inhalte und wähle Dokument PDF hochladen. Du kannst eine PDF-Datei oder ein Bild verwenden.",
)
add_step(
    doc,
    "2",
    "Art und Freigabe wählen",
    "Vergib einen verständlichen Namen und eine Art, zum Beispiel Speiseplan oder Elterninformation. Markiere interne Dokumente als nur intern.",
)
add_step(
    doc,
    "3",
    "Vorschau prüfen",
    "Kontrolliere alle Seiten. Sehr kleine Schrift ist auf einem Bilderrahmen oft schwer zu lesen.",
)
add_step(
    doc,
    "4",
    "Dokument zuordnen",
    "Wähle das Dokument anschließend beim gewünschten Bilderrahmen, in einem Zeitfenster oder in der gemeinsamen Dokumentenliste aus.",
)

doc.add_heading("Gemeinsame Dokumentenliste", level=2)
p = doc.add_paragraph(
    "Mit der gemeinsamen Dokumentenliste kannst du ein Dokument mehreren Gruppen bereitstellen, ohne jeden Rahmen einzeln zu ändern."
)
add_bullet(doc, "Alle Gruppen: Das Dokument gilt auch für später neu angelegte Gruppen.")
add_bullet(doc, "Ausgewählte Gruppen: Das Dokument gilt nur für die angehakten Gruppen.")
add_bullet(doc, "Keine: Das Dokument bleibt gespeichert, erscheint aber nicht mehr in der gemeinsamen Liste.")
add_tip(
    doc,
    "Wichtig",
    "Ein hochgeladenes Dokument wird nicht automatisch angezeigt. Es muss einer Anzeige, einem Zeitfenster oder der gemeinsamen Dokumentenliste zugeordnet werden.",
)

# Anzeige wechseln
doc.add_heading("Anzeige vorübergehend wechseln", level=1)
add_step(
    doc,
    "1",
    "Ziel wählen",
    "Öffne Bilderrahmen und wähle Anzeige wechseln. Wähle ein Gerät, eine Gruppe oder alle Geräte.",
)
add_step(
    doc,
    "2",
    "Inhalt wählen",
    "Wähle Fotos, Dokument, Meldungen oder Gemischt. Bei geteilten Bildschirmen kannst du Hauptanzeige oder zusätzliche Anzeige wählen.",
)
add_step(
    doc,
    "3",
    "Ende festlegen",
    "Mit einem Ende kehrt der Rahmen danach automatisch zum normalen Ablauf zurück. Ohne Ende bleibt die Auswahl bestehen.",
)
add_tip(
    doc,
    "Zurück zur normalen Anzeige",
    "Wähle Automatisch, wenn der Bilderrahmen wieder seine gespeicherten Regeln verwenden soll.",
)

# Rahmen
doc.add_heading("Bilderrahmen einstellen", level=1)
p = doc.add_paragraph(
    "Öffne Bilderrahmen und wähle beim gewünschten Gerät Verwalten. Für den Alltag sind vor allem diese Einstellungen wichtig:"
)
p.paragraph_format.keep_with_next = True
add_bullet(doc, "Standort und Freigaben: Name, Gruppe und Zulassung interner Inhalte")
add_bullet(doc, "Hauptanzeige Standardinhalt: Fotos, Dokument, Meldungen oder eine gemischte Wiedergabe")
add_bullet(doc, "Bild und Darstellung: Ausrichtung, Bildausschnitt, Reihenfolge und Übergänge")
add_bullet(doc, "Betriebszeiten: Gemeinsame Zeiten verwenden oder für diesen Rahmen eigene Zeiten festlegen")
add_bullet(doc, "Gerät und Wartung: Vorschau öffnen, Abgleich anstoßen oder den Player neu starten")

doc.add_heading("Zwei Inhalte gleichzeitig", level=2)
p = doc.add_paragraph(
    "Unter Bildschirm teilen kannst du zwei Bereiche einschalten. Zum Beispiel laufen links Fotos und rechts Meldungen oder der Speiseplan."
)
add_bullet(doc, "Lege fest, ob die Bereiche nebeneinander oder untereinander erscheinen.")
add_bullet(doc, "Wähle die Größe und die Seite der Hauptanzeige.")
add_bullet(doc, "Lege für den zweiten Bereich einen eigenen Standardinhalt fest.")
add_tip(
    doc,
    "Prüfen",
    "Nutze die Vorschau und kontrolliere die echte Anzeige am Standort. Dokumente können in einem kleinen Bereich zu schwer lesbar sein.",
)

# Zeiten
doc.add_heading("Anzeigezeiten festlegen", level=1)
doc.add_heading("Betriebszeiten", level=2)
p = doc.add_paragraph("Unter Betriebszeiten legst du fest, wann die Bilderrahmen anzeigen dürfen.")
add_bullet(doc, "Begrenzung aus: Die Rahmen dürfen rund um die Uhr anzeigen.")
add_bullet(doc, "Begrenzung an: Die Rahmen zeigen nur an den gewählten Tagen und Uhrzeiten.")
add_bullet(doc, "Begrenzung an, aber kein Zeitfenster: Die Rahmen bleiben immer in Ruhe.")
add_tip(
    doc,
    "Eigene Zeiten",
    "Ein einzelner Rahmen kann unter Verwalten eigene Betriebszeiten erhalten. Diese ersetzen den gemeinsamen Zeitplan vollständig.",
)

doc.add_heading("Inhalte zu bestimmten Zeiten", level=2)
p = doc.add_paragraph(
    "Mit Inhaltszeitfenstern wechselst du nur den Inhalt, zum Beispiel werktags von 11 bis 13 Uhr zum Speiseplan. Die Betriebszeiten werden dadurch nicht verlängert."
)
add_step(doc, "1", "Funktion einschalten", "Aktiviere beim Bilderrahmen Inhalte zeitabhängig wechseln.")
add_step(doc, "2", "Zeit und Inhalt wählen", "Wähle Bereich, Wochentage, Uhrzeit und gewünschten Inhalt.")
add_step(
    doc,
    "3",
    "Reihenfolge prüfen",
    "Wenn sich Zeitfenster überschneiden, gilt pro Bereich das oberste passende Fenster. Mit den Pfeilen änderst du die Reihenfolge.",
)

doc.add_heading("Welche Einstellung gewinnt", level=1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = (Cm(1.3), Cm(5.4), Cm(9.5))
mark_header_row(table.rows[0])
for i, (text, width) in enumerate(zip(("Rang", "Einstellung", "Bedeutung"), widths)):
    table.columns[i].width = width
    cell = table.rows[0].cells[i]
    cell.text = text
    set_cell_shading(cell, GREEN)
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
rows = (
    ("1", "Sofortmeldung", "Erscheint immer im Vollbild"),
    ("2", "Betriebszeit oder Bewegungspause", "Der ganze Bildschirm bleibt in Ruhe"),
    ("3", "Manuelle Auswahl", "Gilt bis zum Ende oder bis Automatisch gewählt wird"),
    ("4", "Inhaltszeitfenster", "Bestimmt den Inhalt zur passenden Zeit"),
    ("5", "Standardinhalt", "Gilt, wenn keine höhere Regel greift"),
)
for row_no, values in enumerate(rows):
    cells = table.add_row().cells
    for i, value in enumerate(values):
        cells[i].text = value
        cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if row_no % 2 == 1:
            set_cell_shading(cells[i], PALE_GREEN)
        if i == 0:
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for row in table.rows:
    for cell in row.cells:
        set_cell_margins(cell)
set_table_borders(table)

# Status
doc.add_heading("Status prüfen", level=1)
p = doc.add_paragraph("Unter Bilderrahmen helfen dir diese Angaben bei der Kontrolle:")
p.paragraph_format.keep_with_next = True
add_bullet(doc, "Verbunden: Der Rahmen hat sich kürzlich gemeldet.")
add_bullet(doc, "Soll Anzeige: Das soll der Rahmen nach den aktuellen Regeln zeigen.")
add_bullet(
    doc,
    "Gemeldet: Das hat der Rahmen zuletzt selbst gemeldet. Bei einem getrennten Gerät kann der Eintrag älter sein.",
)
add_bullet(
    doc, "Letzter Abgleich: Zeitpunkt der letzten erfolgreichen Übernahme von Einstellungen und Inhalten."
)
add_bullet(doc, "NAS Status: Zeigt bei jeder Gruppe, ob der Fotoordner erfolgreich gelesen wurde.")
add_bullet(doc, "Vorschau: Zeigt die vorgesehenen Inhalte. Sie ist kein Live Foto des Bildschirms.")

# Fehlerhilfe
doc.add_heading("Wenn etwas nicht klappt", level=1)
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Cm(6.0)
table.columns[1].width = Cm(10.2)
mark_header_row(table.rows[0])
for i, text in enumerate(("Problem", "Das kannst du tun")):
    cell = table.rows[0].cells[i]
    cell.text = text
    set_cell_shading(cell, GREEN)
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
problems = (
    (
        "Neue Fotos fehlen",
        "Kurz warten. Prüfen, ob die Fotos im richtigen Gruppenordner liegen und vollständig kopiert wurden. Danach NAS abgleichen wählen.",
    ),
    (
        "Ein Dokument erscheint nicht",
        "Prüfen, ob es einer Anzeige, einem Zeitfenster oder der gemeinsamen Dokumentenliste zugeordnet wurde.",
    ),
    (
        "Der Rahmen zeigt den falschen Inhalt",
        "Soll Anzeige prüfen. Danach kontrollieren, ob eine manuelle Auswahl, ein Zeitfenster oder eine Sofortmeldung Vorrang hat.",
    ),
    (
        "Der Bildschirm ist schwarz",
        "Betriebszeiten und gegebenenfalls Bewegungssensor prüfen. Eine Ruhezeit kann beabsichtigt sein.",
    ),
    (
        "Ein Gerät ist nicht verbunden",
        "Strom und Netzwerk am Standort prüfen. Danach Abgleichen wählen und kurz warten.",
    ),
    (
        "Eine Änderung kommt nicht an",
        "Letzter Abgleich prüfen. Bei Bedarf Abgleichen auslösen. Player neu starten startet nur die Anzeige Anwendung, nicht das ganze Gerät.",
    ),
)
for row_no, (problem, action) in enumerate(problems):
    cells = table.add_row().cells
    cells[0].text = problem
    cells[1].text = action
    cells[0].paragraphs[0].runs[0].bold = True
    if row_no % 2 == 1:
        for c in cells:
            set_cell_shading(c, PALE_GREEN)
    for c in cells:
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
for row in table.rows:
    for cell in row.cells:
        set_cell_margins(cell, top=125, bottom=125)
set_table_borders(table)

doc.add_heading("Sicher im Alltag", level=1)
add_bullet(doc, "Interne Inhalte nur auf dafür freigegebenen Personalrahmen anzeigen.")
add_bullet(doc, "Vor dem Veröffentlichen Zielgruppe, Inhalt und Endzeit prüfen.")
add_bullet(doc, "Sofortmeldungen nur einsetzen, wenn andere Inhalte wirklich unterbrochen werden sollen.")
add_bullet(doc, "Bei wichtigen Dokumenten die Lesbarkeit am tatsächlichen Bildschirm kontrollieren.")

add_tip(
    doc,
    "Noch unsicher",
    "Ändere zunächst nur ein Gerät oder eine Gruppe und prüfe das Ergebnis in der Vorschau. Danach kannst du die Einstellung auf weitere Rahmen übertragen.",
)

doc.core_properties.title = "KitaFenster Bedienungsanleitung"
doc.core_properties.subject = "Kurzanleitung für Anwenderinnen und Anwender"
doc.core_properties.author = "KitaFenster"
doc.core_properties.keywords = "KitaFenster, Bedienung, Anleitung, Bilderrahmen"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
